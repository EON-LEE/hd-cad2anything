import pandas as pd
from openpyxl import load_workbook
from shapely.geometry import box

from cad2anything.generators import GeneratorContext, run_generator
from cad2anything.model import Entity, EntityKind, Floor, Revision, SpatialModel


def _model() -> SpatialModel:
    model = SpatialModel(revision=Revision("R1", effective_date="2025-03-01"), floors=[Floor("3F", 2)])
    units = [
        ("3F-N01", "BR0001", "메종 루미", box(0, 0, 10, 10)),
        ("3F-N02", "BR0002", "아틀리에 노아", box(10, 0, 20, 10)),
        ("3F-N03", None, None, box(20, 0, 30, 10)),
    ]
    for code, brand_id, brand_name, geom in units:
        model.add(
            Entity(
                f"U-{code}",
                EntityKind.UNIT,
                "3F",
                geom,
                attributes={
                    "unit_code": code,
                    "brand_id": brand_id,
                    "brand_name": brand_name,
                    "category": "WOMEN_CONTEMPORARY",
                    "vacant": brand_id is None,
                },
            )
        )
    return model


def _data() -> dict[str, pd.DataFrame]:
    tenancies = pd.DataFrame(
        [
            ["BR0001", "3F-N01", "R1", "2025-03-01", "2025-08-31", 0.33],
            ["BR0002", "3F-N02", "R1", "2025-03-01", "2025-08-31", 0.33],
        ],
        columns=["brand_id", "unit_code", "rev_id", "valid_from", "valid_to", "commission_rate"],
    )
    terminals = pd.DataFrame(
        [
            ["P001", "BR0001", "3F-N01", "2024-10-01", "2026-09-30"],
            # The brand moved into 3F-N02 for R1; this stale row was not updated.
            ["P002", "BR0002", "3F-N03", "2024-10-01", "2026-09-30"],
        ],
        columns=["pos_id", "brand_id", "unit_code", "valid_from", "valid_to"],
    )
    transactions = pd.DataFrame(
        [
            ["T1", "2025-03-02 11:00:00", "P001", 100_000, 50.0],
            ["T2", "2025-03-03 12:00:00", "P001", -20_000, 50.0],
            ["T3", "2025-03-03 12:00:00", "P002", 70_000, 50.0],
            ["T4", "2024-12-01 12:00:00", "P001", 90_000, 50.0],
        ],
        columns=["txn_id", "ts", "pos_id", "amount_krw", "weight"],
    )
    return {"tenancies": tenancies, "pos_terminals": terminals, "transactions": transactions}


def test_pos_space_kpi_maps_terminals_and_flags_drift(tmp_path):
    record = run_generator("pos_space_kpi", GeneratorContext(model=_model(), out_dir=tmp_path, data=_data()))

    assert record.status == "qa_failed"
    messages = " ".join(issue["message"] for issue in record.issues)
    assert "P002" in messages and "3F-N03" in messages
    assert "3F-N02" in messages  # occupied unit without a terminal
    assert record.summary["window"] == ["2025-03-01", "2025-08-31"]

    sheet = load_workbook(tmp_path / "R1" / "pos_space_kpi" / "pos_space_kpi.xlsx")["매장KPI"]
    rows = {row[1]: row for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert rows["3F-N01"][5] == (100_000 - 20_000) * 50
    assert rows["3F-N01"][6] == 100
    assert rows["3F-N03"][5] == 70_000 * 50


def test_pos_space_kpi_requires_tenancies_for_revision(tmp_path):
    data = _data()
    data["tenancies"] = data["tenancies"].assign(rev_id="R0")

    record = run_generator("pos_space_kpi", GeneratorContext(model=_model(), out_dir=tmp_path, data=data))

    assert record.status == "qa_failed" and record.files == []
