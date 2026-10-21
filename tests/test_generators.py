import csv
from pathlib import Path

import pytest
from openpyxl import load_workbook
from shapely.geometry import LineString, Polygon, box

from cad2anything.generators import (
    GeneratorContext,
    approve,
    load_specs,
    plan_impact,
    reject,
    run_generator,
)
from cad2anything.model import (
    ChangeItem,
    ChangeSet,
    ChangeType,
    Entity,
    EntityKind,
    Floor,
    Revision,
    SpatialModel,
)


def _small_model(rev_id: str = "R0") -> SpatialModel:
    model = SpatialModel(revision=Revision(rev_id), floors=[Floor("3F", 2, "여성 컨템포러리")])
    model.add(Entity("X-3F", EntityKind.EXTERIOR, "3F", box(0, 0, 40, 20)))
    model.add(Entity("W-3F-1", EntityKind.WALKWAY, "3F", box(0, 8, 40, 12)))
    model.add(
        Entity(
            "U-3F-N01",
            EntityKind.UNIT,
            "3F",
            box(0, 12, 20, 20),
            attributes={
                "unit_code": "3F-N01",
                "brand_id": "BR0001",
                "brand_name": "메종 루미",
                "category": "WOMEN_CONTEMPORARY",
            },
        )
    )
    model.add(
        Entity(
            "U-3F-S01",
            EntityKind.UNIT,
            "3F",
            box(0, 0, 20, 8),
            attributes={"unit_code": "3F-S01", "vacant": True},
        )
    )
    model.add(
        Entity(
            "V-3F-ESC-A",
            EntityKind.VERTICAL_LINK,
            "3F",
            box(30, 14, 34, 18),
            attributes={"link_id": "ESC-A", "link_type": "escalator"},
        )
    )
    model.add(
        Entity(
            "A-3F-RESTROOM-1",
            EntityKind.AMENITY,
            "3F",
            box(30, 0, 36, 6),
            attributes={"amenity_type": "RESTROOM"},
        )
    )
    model.add(
        Entity(
            "O-3F-ENT",
            EntityKind.OPENING,
            "3F",
            LineString([(18, 0), (22, 0)]),
            attributes={"opening_id": "ENT"},
        )
    )
    for index, (x, y) in enumerate([(3, 15), (8, 15)], start=1):
        model.add(
            Entity(
                f"F-3F-N01-F00{index}",
                EntityKind.FIXTURE,
                "3F",
                Polygon(box(x, y, x + 1.2, y + 0.5)),
                attributes={
                    "fixture_id": f"3F-N01-F00{index}",
                    "fixture_type": "HANGER",
                    "unit_code": "3F-N01",
                },
            )
        )
    return model


def test_specs_load_with_valid_triggers():
    specs = load_specs()
    assert {"fixture_schedule", "area_schedule", "floor_directory_svg", "pos_space_kpi"} <= set(specs)
    assert specs["area_schedule"].risk_tier == "high"


def test_impact_plan_is_selective():
    specs = load_specs()
    fixture_only = ChangeSet(
        "R0", "R1", [ChangeItem(ChangeType.FIXTURE_CHANGE, "B1", ["B1-C01"], ["B1-C01"])]
    )
    walkway_only = ChangeSet("R0", "R1", [ChangeItem(ChangeType.WALKWAY_CHANGE, "5F")])

    assert [entry.generator for entry in plan_impact(fixture_only, specs)] == ["fixture_schedule"]
    walkway_plan = [entry.generator for entry in plan_impact(walkway_only, specs)]
    assert walkway_plan == ["area_schedule", "floor_directory_svg"]


def test_fixture_and_area_schedules(tmp_path: Path):
    model = _small_model()
    fixtures = run_generator("fixture_schedule", GeneratorContext(model=model, out_dir=tmp_path))
    areas = run_generator("area_schedule", GeneratorContext(model=model, out_dir=tmp_path))

    assert fixtures.summary["fixtures"] == 2 and fixtures.status == "awaiting_approval"
    workbook = load_workbook(tmp_path / "R0" / "fixture_schedule" / "fixture_schedule.xlsx")
    assert workbook.sheetnames == ["집기목록", "매장별요약"]

    with (tmp_path / "R0" / "area_schedule" / "area_schedule.csv").open(encoding="utf-8-sig") as handle:
        rows = list(csv.DictReader(handle))
    lease_total = sum(float(row["임대면적(㎡)"]) for row in rows)
    assert lease_total == pytest.approx(800.0, rel=1e-3)
    assert areas.summary["exclusive_m2"] == pytest.approx(320.0)
    assert not [issue for issue in areas.issues if issue["severity"] == "error"]


def test_floor_directory_svg_labels_units(tmp_path: Path):
    record = run_generator("floor_directory_svg", GeneratorContext(model=_small_model(), out_dir=tmp_path))

    svg = (tmp_path / "R0" / "floor_directory_svg" / "floor_directory_3F.svg").read_text(encoding="utf-8")
    assert record.status == "published"
    assert "메종 루미" in svg and "공실" in svg and "ESC-A" in svg and "화장실" in svg


def test_runs_are_idempotent_and_reviewable(tmp_path: Path):
    model = _small_model()
    first = run_generator("fixture_schedule", GeneratorContext(model=model, out_dir=tmp_path))
    second = run_generator("fixture_schedule", GeneratorContext(model=model, out_dir=tmp_path))
    assert first.job_id == second.job_id and first.created_at == second.created_at

    record_path = tmp_path / "R0" / "fixture_schedule" / "artifact.json"
    approved = approve(record_path, reviewer="space-planner", note="확인")
    assert approved.status == "published"
    with pytest.raises(ValueError):
        reject(record_path, reviewer="space-planner")

    changed = _small_model()
    changed.entities["U-3F-N01"].attributes["brand_name"] = "아틀리에 노아"
    third = run_generator("fixture_schedule", GeneratorContext(model=changed, out_dir=tmp_path))
    assert third.job_id != first.job_id and third.status == "awaiting_approval"


def test_fixture_outside_unit_is_an_error(tmp_path: Path):
    model = _small_model()
    model.entities["F-3F-N01-F001"].attributes["unit_code"] = "3F-N99"

    record = run_generator("fixture_schedule", GeneratorContext(model=model, out_dir=tmp_path))

    assert record.status == "qa_failed"
    assert any("3F-N01-F001" in issue["message"] for issue in record.issues)
